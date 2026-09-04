from pathlib import Path
import json
import re
import time
import unicodedata

import pandas as pd

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# ============================================================
# PARAMÈTRES
# ============================================================

BASE_DIR = Path(__file__).resolve().parents[1]

LLM_DIR = BASE_DIR / "LLM_Tagged_Database"
CLEANED_DIR = BASE_DIR / "Cleaned_Database"

OUTPUT_DIR = BASE_DIR / "outputs_graphs" / "couts_dataset"
OUTPUT_EXCEL = OUTPUT_DIR / "couts_dataset.xlsx"

FICHES_ACTIONS_KEYWORD = "Fiches_actions_database"

# Plus ce seuil est bas, plus le script accepte de résultats,
# mais avec davantage de risques de faux positifs.
MINIMUM_CONFIDENCE_SCORE = 55

# Nombre de caractères conservés avant et après le montant
# dans l'extrait du texte brut.
EXCERPT_WINDOW = 350


# ============================================================
# CHARGEMENT ET NORMALISATION
# ============================================================

def load_json(path):
    with open(path, "r", encoding="utf-8") as file:
        return json.load(file)


def normalize_text(value):
    """
    Transforme une valeur JSON en texte simple.
    """

    if value is None:
        return ""

    if isinstance(value, list):
        return " ".join(
            normalize_text(item)
            for item in value
        )

    if isinstance(value, dict):
        return " ".join(
            normalize_text(item)
            for item in value.values()
        )

    text = str(value)

    text = text.replace("\u00a0", " ")
    text = text.replace("\u202f", " ")
    text = text.replace("\r\n", "\n")
    text = text.replace("\r", "\n")

    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n+", "\n", text)

    return text.strip()


def normalize_filename(value):
    """
    Normalise un nom de fichier afin de faciliter le rapprochement
    entre un fichier JSON et son fichier texte.
    """

    value = str(value).lower()

    value = unicodedata.normalize("NFD", value)
    value = "".join(
        character
        for character in value
        if unicodedata.category(character) != "Mn"
    )

    value = re.sub(r"[^a-z0-9]+", "", value)

    return value


def extract_values(value):
    if value is None:
        return []

    if isinstance(value, list):
        results = []

        for item in value:
            results.extend(extract_values(item))

        return results

    if isinstance(value, dict):
        results = []

        for item in value.values():
            results.extend(extract_values(item))

        return results

    cleaned = normalize_text(value)

    return [cleaned] if cleaned else []


def get_nested(data, *keys):
    current = data

    for key in keys:
        if not isinstance(current, dict):
            return None

        current = current.get(key)

    return current


# ============================================================
# IDENTIFICATION DES FICHES
# ============================================================

def is_fiche_action(json_path):
    return (
        FICHES_ACTIONS_KEYWORD.lower()
        in str(json_path).lower()
    )


def is_habitat_logement(data):
    thematiques = extract_values(data.get("thematique"))

    joined = " ".join(
        thematique.lower()
        for thematique in thematiques
    )

    keywords = [
        "habitat",
        "logement",
        "habitat et logement",
        "adaptation du logement",
        "adaptation des logements",
        "résidence",
        "residence",
        "domicile",
        "hébergement",
        "hebergement",
        "bailleur",
        "copropriété",
        "copropriete"
    ]

    return any(
        keyword in joined
        for keyword in keywords
    )


# ============================================================
# INDEXATION DES TEXTES BRUTS
# ============================================================

def build_cleaned_text_index():
    """
    Construit un index des fichiers TXT une seule fois.

    Cela évite de parcourir entièrement Cleaned_Database pour chaque JSON.
    """

    index = {}

    if not CLEANED_DIR.exists():
        print(
            f"Attention : le dossier de textes bruts n'existe pas : "
            f"{CLEANED_DIR}"
        )
        return index

    txt_files = list(CLEANED_DIR.rglob("*.txt"))

    print(
        f"Indexation de {len(txt_files)} fichiers texte brut..."
    )

    for txt_path in txt_files:
        normalized_stem = normalize_filename(txt_path.stem)

        if not normalized_stem:
            continue

        index.setdefault(normalized_stem, []).append(txt_path)

    return index


def find_matching_cleaned_file(json_path, cleaned_index):
    """
    Retrouve le texte brut correspondant au JSON.

    La recherche se fait :
    1. avec le même nom de fichier ;
    2. avec un nom normalisé ;
    3. avec une correspondance partielle prudente.
    """

    if not cleaned_index:
        return None, "aucun index de textes bruts"

    normalized_json_stem = normalize_filename(json_path.stem)

    # Correspondance exacte après normalisation
    exact_matches = cleaned_index.get(
        normalized_json_stem,
        []
    )

    if len(exact_matches) == 1:
        return exact_matches[0], "nom exact"

    if len(exact_matches) > 1:
        # Si plusieurs résultats, on privilégie un chemin
        # contenant le même dossier parent.
        json_parent = normalize_filename(
            json_path.parent.name
        )

        for candidate in exact_matches:
            if json_parent in normalize_filename(
                candidate.parent.name
            ):
                return candidate, "nom exact et dossier similaire"

        return exact_matches[0], "nom exact, plusieurs fichiers possibles"

    # Correspondance partielle
    partial_matches = []

    if len(normalized_json_stem) >= 12:
        for normalized_txt_stem, paths in cleaned_index.items():
            if (
                normalized_json_stem in normalized_txt_stem
                or normalized_txt_stem in normalized_json_stem
            ):
                partial_matches.extend(paths)

    if len(partial_matches) == 1:
        return partial_matches[0], "nom partiellement similaire"

    return None, "aucun texte brut correspondant"


def read_matching_cleaned_text(json_path, cleaned_index):
    txt_path, matching_status = find_matching_cleaned_file(
        json_path,
        cleaned_index
    )

    if txt_path is None:
        return "", "", matching_status

    try:
        text = txt_path.read_text(
            encoding="utf-8",
            errors="ignore"
        )

        return text, str(txt_path), matching_status

    except Exception as error:
        return (
            "",
            str(txt_path),
            f"erreur de lecture : {error}"
        )


# ============================================================
# CONSTRUCTION DES TEXTES À ANALYSER
# ============================================================

def build_search_contents(
    data,
    json_path,
    cleaned_index
):
    """
    Conserve séparément :
    - le texte brut ;
    - les champs JSON importants ;
    - le JSON complet.
    """

    (
        raw_text,
        raw_text_path,
        raw_matching_status
    ) = read_matching_cleaned_text(
        json_path,
        cleaned_index
    )

    description = data.get("description", {})

    if not isinstance(description, dict):
        description = {}

    important_values = [
        data.get("titre"),
        data.get("territoire"),
        data.get("echelle"),
        data.get("public_vise"),
        data.get("nature_initiative"),
        data.get("porteur_initiative"),
        data.get("thematique"),
        data.get("source_site"),
        data.get("statut_action"),

        # Champs financiers éventuellement déjà présents
        data.get("cout"),
        data.get("coût"),
        data.get("cout_initiative"),
        data.get("budget"),
        data.get("financement"),
        data.get("montant"),

        # Description imbriquée
        description.get("description_generale"),
        description.get("contexte"),
        description.get("problematique"),
        description.get("solution_envisagee"),
        description.get("objectifs"),

        # Champs financiers imbriqués éventuels
        description.get("cout"),
        description.get("coût"),
        description.get("budget"),
        description.get("financement"),
        description.get("montant")
    ]

    return {
        "raw_text": raw_text,
        "raw_text_path": raw_text_path,
        "raw_matching_status": raw_matching_status,
        "important_json_text": normalize_text(
            important_values
        ),
        "json_text": normalize_text(data)
    }


# ============================================================
# EXTRACTION D'UN PASSAGE DU TEXTE BRUT
# ============================================================

def extract_raw_excerpt(
    raw_text,
    match_start,
    match_end,
    window=EXCERPT_WINDOW
):
    """
    Extrait un passage du texte brut autour du montant,
    en essayant de conserver des phrases complètes.
    """

    if not raw_text:
        return ""

    start = max(0, match_start - window)
    end = min(
        len(raw_text),
        match_end + window
    )

    excerpt = raw_text[start:end]

    # Recherche d'une limite naturelle avant le montant
    relative_match_start = match_start - start

    left_part = excerpt[:relative_match_start]

    left_boundaries = [
        left_part.rfind("\n"),
        left_part.rfind(". "),
        left_part.rfind("! "),
        left_part.rfind("? "),
        left_part.rfind("; ")
    ]

    best_left_boundary = max(left_boundaries)

    if best_left_boundary >= 0:
        excerpt = excerpt[
            best_left_boundary + 1:
        ]

    # Recherche d'une limite naturelle après le montant
    new_relative_match_end = (
        match_end - start - best_left_boundary - 1
        if best_left_boundary >= 0
        else match_end - start
    )

    right_part = excerpt[
        new_relative_match_end:
    ]

    right_boundaries = [
        position
        for position in [
            right_part.find("\n"),
            right_part.find(". "),
            right_part.find("! "),
            right_part.find("? "),
            right_part.find("; ")
        ]
        if position >= 0
    ]

    if right_boundaries:
        best_right_boundary = min(right_boundaries)

        excerpt = excerpt[
            :new_relative_match_end
            + best_right_boundary
            + 1
        ]

    return normalize_text(excerpt)


# ============================================================
# NORMALISATION ET ANALYSE DES MONTANTS
# ============================================================

def normalize_amount_for_comparison(amount):
    """
    Normalise un montant pour repérer les doublons.
    """

    normalized = amount.lower()
    normalized = normalized.replace("\u00a0", "")
    normalized = normalized.replace("\u202f", "")
    normalized = normalized.replace(" ", "")
    normalized = normalized.replace(".", "")
    normalized = normalized.replace(",", ".")

    normalized = normalized.replace("euros", "€")
    normalized = normalized.replace("euro", "€")
    normalized = normalized.replace("eur", "€")

    return normalized


def parse_numeric_amount(amount):
    """
    Convertit approximativement un montant en valeur numérique
    pour appliquer certains contrôles.

    Retourne None pour les fourchettes.
    """

    amount_lower = amount.lower()

    if (
        "entre" in amount_lower
        or re.search(r"\bà\b", amount_lower)
    ):
        return None

    multiplier = 1

    if re.search(r"\bk\s*€", amount_lower):
        multiplier = 1_000

    elif re.search(r"\bm\s*€", amount_lower):
        multiplier = 1_000_000

    number_match = re.search(
        r"\d[\d\s\u00a0\u202f.]*"
        r"(?:,\d+)?",
        amount_lower
    )

    if not number_match:
        return None

    number = number_match.group(0)

    number = number.replace("\u00a0", "")
    number = number.replace("\u202f", "")
    number = number.replace(" ", "")

    # Un point placé entre groupes de 3 chiffres
    # est considéré comme séparateur de milliers.
    if re.fullmatch(
        r"\d{1,3}(?:\.\d{3})+",
        number
    ):
        number = number.replace(".", "")

    number = number.replace(",", ".")

    try:
        return float(number) * multiplier
    except ValueError:
        return None


def detect_cost(search_contents):
    """
    Détecte le coût le plus probable d'une initiative.

    Retourne :
    - coût retenu ;
    - statut de détection ;
    - extrait du texte brut ;
    - contexte de détection ;
    - source de la détection ;
    - score ;
    - autres montants détectés.
    """

    raw_text = search_contents.get(
        "raw_text",
        ""
    )

    important_json_text = search_contents.get(
        "important_json_text",
        ""
    )

    json_text = search_contents.get(
        "json_text",
        ""
    )

    candidates = []

    # --------------------------------------------------------
    # Expressions monétaires explicites
    # --------------------------------------------------------

    number_pattern = (
        r"(?:"
        r"\d{1,3}(?:[\s\u00a0\u202f.]\d{3})+"
        r"|\d+"
        r")"
        r"(?:[,.]\d+)?"
    )

    currency_pattern = (
        r"(?:"
        r"€"
        r"|euros?"
        r"|eur"
        r"|k\s*€"
        r"|m\s*€"
        r")"
    )

    explicit_money_pattern = re.compile(
        rf"""
        # Fourchette : entre 5 000 et 8 000 euros
        (?:
            entre\s+
            {number_pattern}
            \s*
            {currency_pattern}?
            \s+
            (?:et|à)
            \s+
            {number_pattern}
            \s*
            {currency_pattern}
        )

        |

        # Fourchette : de 5 000 à 8 000 euros
        (?:
            de\s+
            {number_pattern}
            \s*
            {currency_pattern}?
            \s+
            à
            \s+
            {number_pattern}
            \s*
            {currency_pattern}
        )

        |

        # Montant classique : 10 000 €, 10000 euros, 15 k€
        (?:
            {number_pattern}
            \s*
            {currency_pattern}
        )

        |

        # Devise avant le montant : EUR 10 000
        (?:
            (?:€|eur)
            \s*
            {number_pattern}
        )
        """,
        flags=re.IGNORECASE | re.VERBOSE
    )

    # --------------------------------------------------------
    # Montants sans devise, uniquement avec contexte explicite
    # --------------------------------------------------------

    contextual_amount_pattern = re.compile(
        rf"""
        (?:
            coût
            |cout
            |budget
            |montant
            |enveloppe
        )
        (?:
            \s+
            (?:
                total
                |global
                |prévisionnel
                |previsionnel
                |estimé
                |estime
                |de\s+l'opération
                |de\s+l’operation
                |du\s+projet
                |de\s+l'initiative
                |de\s+l’initiative
            )
        )?
        \s*
        (?:
            est
            |était
            |etait
            |s'élève\s+à
            |s’élève\s+à
            |s'élevait\s+à
            |s’élevait\s+à
            |représente
            |représentait
            |de
            |:
            |=
        )?
        \s*
        (?P<amount>
            \d{{1,3}}
            (?:[\s\u00a0\u202f.]\d{{3}})+
            (?:[,.]\d+)?
            |
            \d{{4,}}
            (?:[,.]\d+)?
        )
        """,
        flags=re.IGNORECASE | re.VERBOSE
    )

    positive_keywords = {
        "coût total": 200,
        "cout total": 200,
        "budget total": 200,
        "montant total": 195,
        "coût global": 190,
        "cout global": 190,
        "budget global": 190,
        "montant global": 185,
        "budget prévisionnel": 180,
        "budget previsionnel": 180,
        "coût prévisionnel": 180,
        "cout previsionnel": 180,
        "enveloppe globale": 175,
        "enveloppe budgétaire": 175,
        "enveloppe budgetaire": 175,
        "coût de l'opération": 170,
        "coût de l’opération": 170,
        "cout de l'operation": 170,
        "coût du projet": 170,
        "cout du projet": 170,
        "budget du projet": 165,
        "coût de l'initiative": 165,
        "coût de l’initiative": 165,
        "cout de l'initiative": 165,
        "au total": 140,
        "s'élève à": 130,
        "s’élève à": 130,
        "coût": 100,
        "cout": 100,
        "budget": 95,
        "montant": 90,
        "enveloppe": 85,
        "financement": 70,
        "financé": 65,
        "financée": 65,
        "subvention": 55,
        "dépense": 50,
        "dépenses": 50,
        "investissement": 50,
        "reste à charge": 35,
        "participation financière": 30,
        "tarif": 20,
        "prix": 20
    }

    negative_keywords = {
        "par personne": -90,
        "par participant": -90,
        "par bénéficiaire": -90,
        "par bénéficiaire et par an": -100,
        "par ménage": -65,
        "par logement": -35,
        "par mois": -75,
        "mensuel": -70,
        "par jour": -65,
        "par heure": -65,
        "de l'heure": -65,
        "prix unitaire": -75,
        "tarif unitaire": -75,
        "plafond de ressources": -120,
        "plafond": -45,
        "dans la limite de": -40,
        "jusqu'à": -25,
        "jusqu’à": -25,
        "aide maximale": -45,
        "subvention maximale": -45,
        "revenu": -110,
        "revenus": -110,
        "salaire": -110,
        "allocation": -90,
        "pension": -90,
        "loyer": -65,
        "loyers": -65,
        "facture moyenne": -55,
        "valeur du bien": -80,
        "prix du logement": -80,
        "mètres carrés": -120,
        "mètre carré": -120,
        "m²": -120,
        "chiffre d'affaires": -120
    }

    strong_total_keywords = [
        "coût total",
        "cout total",
        "budget total",
        "montant total",
        "coût global",
        "cout global",
        "budget global",
        "budget prévisionnel",
        "budget previsionnel",
        "coût prévisionnel",
        "cout previsionnel",
        "coût du projet",
        "cout du projet",
        "coût de l'opération",
        "coût de l’opération",
        "budget du projet",
        "coût de l'initiative",
        "coût de l’initiative",
        "au total"
    ]

    # --------------------------------------------------------
    # Calcul du score d'un candidat
    # --------------------------------------------------------

    def score_candidate(
        amount,
        full_text,
        start,
        end,
        source,
        implicit_currency=False
    ):
        context_start = max(0, start - 350)
        context_end = min(
            len(full_text),
            end + 350
        )

        context = full_text[
            context_start:context_end
        ]

        context_normalized = normalize_text(
            context
        )

        context_lower = context_normalized.lower()

        score = 0

        for keyword, points in positive_keywords.items():
            if keyword in context_lower:
                score += points

        for keyword, points in negative_keywords.items():
            if keyword in context_lower:
                score += points

        # Priorité au texte brut
        if source == "texte brut":
            score += 50

        elif source == "JSON prioritaire":
            score += 25

        else:
            score += 10

        amount_lower = amount.lower()

        if any(
            marker in amount_lower
            for marker in [
                "€",
                "euro",
                "eur"
            ]
        ):
            score += 40

        if implicit_currency:
            score -= 20

        # Analyse de la partie placée avant le montant
        before_80 = full_text[
            max(0, start - 80):start
        ].lower()

        before_160 = full_text[
            max(0, start - 160):start
        ].lower()

        if any(
            keyword in before_80
            for keyword in [
                "coût",
                "cout",
                "budget",
                "montant",
                "enveloppe"
            ]
        ):
            score += 80

        elif any(
            keyword in before_160
            for keyword in [
                "coût",
                "cout",
                "budget",
                "montant",
                "enveloppe"
            ]
        ):
            score += 45

        if any(
            keyword in context_lower
            for keyword in strong_total_keywords
        ):
            score += 90

        numeric_amount = parse_numeric_amount(
            amount
        )

        # Les montants extrêmement faibles sont souvent
        # des tarifs individuels plutôt qu'un coût de projet.
        if numeric_amount is not None:
            if numeric_amount < 10:
                score -= 70

            elif numeric_amount < 100:
                score -= 30

            elif numeric_amount >= 1_000:
                score += 15

        return {
            "amount": normalize_text(amount),
            "score": score,
            "context": context_normalized,
            "start": start,
            "end": end,
            "source": source,
            "implicit_currency": implicit_currency
        }

    # --------------------------------------------------------
    # Recherche dans un texte
    # --------------------------------------------------------

    def search_explicit_amounts(
        text,
        source
    ):
        if not text:
            return

        for match in explicit_money_pattern.finditer(text):
            amount = match.group(0)

            candidate = score_candidate(
                amount=amount,
                full_text=text,
                start=match.start(),
                end=match.end(),
                source=source,
                implicit_currency=False
            )

            candidates.append(candidate)

    def search_contextual_amounts(
        text,
        source
    ):
        if not text:
            return

        for match in contextual_amount_pattern.finditer(text):
            amount = match.group("amount")

            candidate = score_candidate(
                amount=amount,
                full_text=text,
                start=match.start("amount"),
                end=match.end("amount"),
                source=source,
                implicit_currency=True
            )

            candidates.append(candidate)

    # --------------------------------------------------------
    # Ordre de recherche
    # --------------------------------------------------------

    search_explicit_amounts(
        raw_text,
        "texte brut"
    )

    search_contextual_amounts(
        raw_text,
        "texte brut"
    )

    search_explicit_amounts(
        important_json_text,
        "JSON prioritaire"
    )

    search_contextual_amounts(
        important_json_text,
        "JSON prioritaire"
    )

    search_explicit_amounts(
        json_text,
        "JSON complet"
    )

    search_contextual_amounts(
        json_text,
        "JSON complet"
    )

    # --------------------------------------------------------
    # Aucun montant trouvé
    # --------------------------------------------------------

    if not candidates:
        combined_text = normalize_text(
            raw_text
            + " "
            + important_json_text
            + " "
            + json_text
        ).lower()

        financial_words = [
            "coût",
            "cout",
            "budget",
            "montant",
            "financement",
            "subvention",
            "dépense",
            "enveloppe",
            "reste à charge"
        ]

        if any(
            word in combined_text
            for word in financial_words
        ):
            return (
                "non connu",
                "contexte financier détecté mais aucun montant fiable",
                "",
                "",
                "",
                "",
                ""
            )

        return (
            "non connu",
            "aucune information de coût détectée",
            "",
            "",
            "",
            "",
            ""
        )

    # --------------------------------------------------------
    # Suppression des doublons
    # --------------------------------------------------------

    unique_candidates = {}

    for candidate in candidates:
        amount_key = normalize_amount_for_comparison(
            candidate["amount"]
        )

        # Un même montant peut apparaître dans le JSON et dans le texte brut.
        # On conserve de préférence sa version issue du texte brut.
        existing_candidate = unique_candidates.get(
            amount_key
        )

        if existing_candidate is None:
            unique_candidates[amount_key] = candidate

        elif candidate["score"] > existing_candidate["score"]:
            unique_candidates[amount_key] = candidate

        elif (
            candidate["source"] == "texte brut"
            and existing_candidate["source"] != "texte brut"
        ):
            unique_candidates[amount_key] = candidate

    candidates = list(
        unique_candidates.values()
    )

    candidates.sort(
        key=lambda item: (
            item["score"],
            item["source"] == "texte brut",
            parse_numeric_amount(
                item["amount"]
            ) or 0
        ),
        reverse=True
    )

    best_candidate = candidates[0]

    # --------------------------------------------------------
    # Rejet si le contexte est trop faible
    # --------------------------------------------------------

    if best_candidate["score"] < MINIMUM_CONFIDENCE_SCORE:
        other_amounts = " ; ".join(
            candidate["amount"]
            for candidate in candidates[:5]
        )

        return (
            "non connu",
            "montant détecté mais contexte insuffisant",
            "",
            best_candidate["context"],
            best_candidate["source"],
            best_candidate["score"],
            other_amounts
        )

    # --------------------------------------------------------
    # Extraction du passage brut
    # --------------------------------------------------------

    raw_excerpt = ""

    if best_candidate["source"] == "texte brut":
        raw_excerpt = extract_raw_excerpt(
            raw_text,
            best_candidate["start"],
            best_candidate["end"]
        )

    elif raw_text:
        # Le candidat vient du JSON : on essaie de retrouver
        # ce même montant dans le texte brut.
        amount_parts = re.findall(
            r"\d+",
            best_candidate["amount"]
        )

        if amount_parts:
            flexible_amount_pattern = (
                r"[\s\u00a0\u202f.,]*".join(
                    re.escape(part)
                    for part in amount_parts
                )
            )

            raw_match = re.search(
                flexible_amount_pattern,
                raw_text,
                flags=re.IGNORECASE
            )

            if raw_match:
                raw_excerpt = extract_raw_excerpt(
                    raw_text,
                    raw_match.start(),
                    raw_match.end()
                )

    # --------------------------------------------------------
    # Autres montants détectés
    # --------------------------------------------------------

    other_amounts = []

    selected_key = normalize_amount_for_comparison(
        best_candidate["amount"]
    )

    for candidate in candidates[1:]:
        candidate_key = normalize_amount_for_comparison(
            candidate["amount"]
        )

        if candidate_key == selected_key:
            continue

        if candidate["amount"] not in other_amounts:
            other_amounts.append(
                candidate["amount"]
            )

        if len(other_amounts) >= 5:
            break

    # --------------------------------------------------------
    # Résultat final
    # --------------------------------------------------------

    if best_candidate["implicit_currency"]:
        selected_amount = (
            best_candidate["amount"]
            + " € ?"
        )

        status = (
            "montant probable détecté sans devise explicite"
        )

    else:
        selected_amount = best_candidate["amount"]

        status = (
            "coût le plus probable sélectionné selon le contexte"
        )

    return (
        selected_amount,
        status,
        raw_excerpt,
        best_candidate["context"],
        best_candidate["source"],
        best_candidate["score"],
        " ; ".join(other_amounts)
    )


# ============================================================
# MISE EN FORME DU FICHIER EXCEL
# ============================================================

def format_excel_workbook(writer):
    workbook = writer.book

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7"
    )

    detected_fill = PatternFill(
        fill_type="solid",
        fgColor="E2F0D9"
    )

    unknown_fill = PatternFill(
        fill_type="solid",
        fgColor="FCE4D6"
    )

    for sheet_name, worksheet in writer.sheets.items():
        worksheet.freeze_panes = "A2"

        if worksheet.max_row >= 1:
            worksheet.auto_filter.ref = (
                worksheet.dimensions
            )

        # En-têtes
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True
            )

        header_positions = {
            cell.value: cell.column
            for cell in worksheet[1]
        }

        # Largeurs spécifiques
        widths = {
            "fichier": 35,
            "titre": 45,
            "type_fiche": 18,
            "thematique": 40,
            "habitat_et_logement": 20,
            "porteur_initiative": 35,
            "cout_initiative": 22,
            "statut_detection": 42,
            "score_detection": 18,
            "source_detection": 20,
            "extrait_texte_brut": 100,
            "contexte_detection": 80,
            "autres_montants_detectes": 40,
            "correspondance_texte_brut": 35,
            "fichier_texte_brut": 60,
            "chemin_fichier": 60
        }

        for column_name, width in widths.items():
            column_index = header_positions.get(
                column_name
            )

            if column_index:
                column_letter = get_column_letter(
                    column_index
                )

                worksheet.column_dimensions[
                    column_letter
                ].width = width

        # Alignement et retour à la ligne
        for row in worksheet.iter_rows(
            min_row=2
        ):
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True
                )

        # Mise en couleur selon le résultat
        cost_column = header_positions.get(
            "cout_initiative"
        )

        if cost_column:
            for row_number in range(
                2,
                worksheet.max_row + 1
            ):
                cost_cell = worksheet.cell(
                    row=row_number,
                    column=cost_column
                )

                if cost_cell.value == "non connu":
                    cost_cell.fill = unknown_fill
                else:
                    cost_cell.fill = detected_fill

        # Hauteur raisonnable des lignes
        for row_number in range(
            2,
            worksheet.max_row + 1
        ):
            worksheet.row_dimensions[
                row_number
            ].height = 60


# ============================================================
# PROGRAMME PRINCIPAL
# ============================================================

def main():
    start_time = time.time()

    print("=" * 75)
    print("RECHERCHE DES COÛTS DANS LE DATASET")
    print("=" * 75)

    print(f"Dossier JSON : {LLM_DIR}")
    print(f"Dossier textes bruts : {CLEANED_DIR}")
    print(f"Fichier Excel : {OUTPUT_EXCEL}")
    print()

    if not LLM_DIR.exists():
        print(
            f"Erreur : le dossier JSON n'existe pas : "
            f"{LLM_DIR}"
        )
        return

    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    cleaned_index = build_cleaned_text_index()

    json_files = list(
        LLM_DIR.rglob("*.json")
    )

    print(
        f"{len(json_files)} fichiers JSON trouvés."
    )
    print()

    rows = []

    error_count = 0
    raw_text_found_count = 0

    for index, json_path in enumerate(
        json_files,
        start=1
    ):
        try:
            data = load_json(json_path)

            if not isinstance(data, dict):
                continue

            search_contents = build_search_contents(
                data,
                json_path,
                cleaned_index
            )

            if search_contents["raw_text"]:
                raw_text_found_count += 1

            (
                cout,
                statut_detection,
                extrait_texte_brut,
                contexte_detection,
                source_detection,
                score_detection,
                autres_montants
            ) = detect_cost(search_contents)

            thematiques = extract_values(
                data.get("thematique")
            )

            rows.append({
                "fichier": json_path.name,

                "titre": normalize_text(
                    data.get("titre")
                ),

                "type_fiche": (
                    "fiche action"
                    if is_fiche_action(json_path)
                    else "autre fiche"
                ),

                "thematique": ", ".join(
                    thematiques
                ),

                "habitat_et_logement": (
                    "oui"
                    if is_habitat_logement(data)
                    else "non"
                ),

                "porteur_initiative": normalize_text(
                    data.get(
                        "porteur_initiative"
                    )
                ),

                "cout_initiative": cout,

                "statut_detection": (
                    statut_detection
                ),

                "score_detection": (
                    score_detection
                ),

                "source_detection": (
                    source_detection
                ),

                "extrait_texte_brut": (
                    extrait_texte_brut
                ),

                "contexte_detection": (
                    contexte_detection
                ),

                "autres_montants_detectes": (
                    autres_montants
                ),

                "correspondance_texte_brut": (
                    search_contents[
                        "raw_matching_status"
                    ]
                ),

                "fichier_texte_brut": (
                    search_contents[
                        "raw_text_path"
                    ]
                ),

                "chemin_fichier": str(
                    json_path
                )
            })

            if index % 100 == 0:
                print(
                    f"{index}/{len(json_files)} "
                    f"fichiers analysés..."
                )

        except Exception as error:
            error_count += 1

            print(
                f"Erreur avec {json_path.name} : "
                f"{error}"
            )

    if not rows:
        print("Aucune fiche exploitable trouvée.")
        return

    df = pd.DataFrame(rows)

    df = df.sort_values(
        by=[
            "habitat_et_logement",
            "cout_initiative",
            "score_detection",
            "titre"
        ],
        ascending=[
            False,
            True,
            False,
            True
        ],
        na_position="last"
    )

    df_habitat = df[
        df["habitat_et_logement"] == "oui"
    ].copy()

    df_couts_detectes = df[
        df["cout_initiative"] != "non connu"
    ].copy()

    df_couts_non_connus = df[
        df["cout_initiative"] == "non connu"
    ].copy()

    df_sans_texte_brut = df[
        df["fichier_texte_brut"] == ""
    ].copy()

    resume = pd.DataFrame([
        {
            "indicateur": "Fichiers JSON trouvés",
            "valeur": len(json_files)
        },
        {
            "indicateur": "Fiches analysées",
            "valeur": len(df)
        },
        {
            "indicateur": "Fiches actions",
            "valeur": len(
                df[
                    df["type_fiche"]
                    == "fiche action"
                ]
            )
        },
        {
            "indicateur": "Autres fiches",
            "valeur": len(
                df[
                    df["type_fiche"]
                    == "autre fiche"
                ]
            )
        },
        {
            "indicateur": "Textes bruts retrouvés",
            "valeur": raw_text_found_count
        },
        {
            "indicateur": "Textes bruts non retrouvés",
            "valeur": len(df_sans_texte_brut)
        },
        {
            "indicateur": "Fiches habitat/logement",
            "valeur": len(df_habitat)
        },
        {
            "indicateur": "Coûts détectés",
            "valeur": len(df_couts_detectes)
        },
        {
            "indicateur": "Coûts non connus",
            "valeur": len(df_couts_non_connus)
        },
        {
            "indicateur": "Erreurs",
            "valeur": error_count
        }
    ])

    with pd.ExcelWriter(
        OUTPUT_EXCEL,
        engine="openpyxl"
    ) as writer:

        df.to_excel(
            writer,
            sheet_name="Tout le dataset",
            index=False
        )

        df_habitat.to_excel(
            writer,
            sheet_name="Habitat logement",
            index=False
        )

        df_couts_detectes.to_excel(
            writer,
            sheet_name="Coûts détectés",
            index=False
        )

        df_couts_non_connus.to_excel(
            writer,
            sheet_name="Coûts non connus",
            index=False
        )

        df_sans_texte_brut.to_excel(
            writer,
            sheet_name="Textes bruts absents",
            index=False
        )

        resume.to_excel(
            writer,
            sheet_name="Résumé",
            index=False
        )

        format_excel_workbook(writer)

    elapsed_time = round(
        time.time() - start_time,
        2
    )

    print()
    print("=" * 75)
    print("EXCEL GÉNÉRÉ AVEC SUCCÈS")
    print("=" * 75)
    print(f"Fichier créé : {OUTPUT_EXCEL}")
    print(f"Fiches analysées : {len(df)}")
    print(
        f"Textes bruts retrouvés : "
        f"{raw_text_found_count}"
    )
    print(
        f"Textes bruts non retrouvés : "
        f"{len(df_sans_texte_brut)}"
    )
    print(
        f"Fiches habitat/logement : "
        f"{len(df_habitat)}"
    )
    print(
        f"Coûts détectés : "
        f"{len(df_couts_detectes)}"
    )
    print(
        f"Coûts non connus : "
        f"{len(df_couts_non_connus)}"
    )
    print(f"Erreurs : {error_count}")
    print(
        f"Temps d'exécution : "
        f"{elapsed_time} secondes"
    )
    print("=" * 75)


if __name__ == "__main__":
    main()