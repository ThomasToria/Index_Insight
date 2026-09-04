"""Evaluate a deterministic rule-based baseline on the LLM reference sample.

The script reuses the completed annotation workbook and evaluates the same gold
labels as the LLM evaluation.  It does not alter the workbook or any JSON record.

For ``public_vise`` and ``echelle``, predictions are derived from the historical
``Final_Tagged_Database`` outputs.  The historical rule-based schema did not
contain ``nature_initiative``.  For that field, the script therefore applies a
documented keyword-only extension to the first 2,500 characters of the cleaned
document: this is the same input limit used by the LLM extraction script and
does not call an LLM.
"""

from __future__ import annotations

import csv
import json
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = PROJECT_ROOT / "Evaluation" / "llm_annotation_template.xlsx"
FINAL_DATABASE = PROJECT_ROOT / "Final_Tagged_Database"
LLM_DATABASE = PROJECT_ROOT / "LLM_Tagged_Database"
OUTPUT_DIR = PROJECT_ROOT / "Evaluation"
BASELINE_RESULTS = OUTPUT_DIR / "rule_based_baseline_results.csv"
COMPARISON_RESULTS = OUTPUT_DIR / "rule_based_vs_llm_comparison.csv"
DIAGNOSTICS = OUTPUT_DIR / "rule_based_baseline_diagnostics.csv"

FIELDS = ("public_vise", "echelle", "nature_initiative")
MAX_CHARS = 2500

# The nature rules are a deterministic extension because the historical
# Final_Tagged_Database has no nature_initiative field.  Their order is part of
# the baseline specification and is intentionally fixed for reproducibility.
NATURE_RULES: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("atelier", ("atelier", "ateliers")),
    ("formation", ("formation", "formations", "former", "formez")),
    ("sensibilisation", ("sensibilisation", "sensibiliser", "campagne de sensibilisation")),
    ("enquête", ("enquête", "questionnaire", "sondage", "étude")),
    ("rapport", ("rapport", "diagnostic", "bilan")),
    ("plaidoyer", ("plaidoyer", "revendication", "interpeller")),
    ("outil numérique", ("application", "plateforme", "site internet", "outil numérique", "carte interactive")),
    ("événement", ("événement", "journée", "forum", "salon", "conférence")),
    ("accompagnement", ("accompagnement", "accompagner", "suivi personnalisé")),
    ("dispositif", ("dispositif", "programme", "expérimentation")),
    ("publication", ("publication", "guide", "livret", "brochure")),
    ("offre de service", ("offre de service", "service proposé", "services proposés")),
    ("groupe de parole", ("groupe de parole", "échange entre pairs")),
    ("consultation citoyenne", ("consultation citoyenne", "concertation", "participation citoyenne")),
    ("lieu / espace dédié", ("lieu dédié", "espace dédié", "tiers-lieu", "local dédié")),
    ("projet d'aménagement", ("aménagement", "réaménagement", "urbanisme", "travaux")),
)


def normalise(value: object) -> str:
    """Normalise accents, case, and whitespace for comparison."""
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKD", text.casefold())
    text = "".join(char for char in text if not unicodedata.combining(char))
    return " ".join(text.split())


def labels(value: object) -> set[str]:
    if isinstance(value, (list, tuple, set)):
        return {normalise(item) for item in value if normalise(item) and normalise(item) != "absent"}
    text = normalise(value)
    if not text or text == "absent":
        return set()
    return {item.strip() for item in text.split("|") if item.strip()}


def safe_ratio(numerator: int, denominator: int) -> float:
    return round(numerator / denominator, 3) if denominator else 0.0


def f1(precision: float, recall: float) -> float:
    return round(2 * precision * recall / (precision + recall), 3) if precision + recall else 0.0


def extract_description(record: dict[str, Any]) -> dict[str, Any]:
    description = record.get("description", {})
    return description if isinstance(description, dict) else {}


def public_from_rule_output(value: object) -> str:
    """Map the historical free-text candidate to the annotation vocabulary."""
    text = normalise(value)
    if not text:
        return "public non connu"

    senior_patterns = (
        "aine", "senior", "retraite", "personne agee", "grand age",
        "vieillissement", "ehpad", "residence autonomie",
    )
    professional_patterns = (
        "professionnel", "intervenant", "animateur", "porteur de projet",
        "porteurs de projet", "agent", "soignant",
    )
    other_patterns = (
        "aidant", "grand public", "enfant", "jeune", "famille", "habitant",
        "usager", "tout public", "personne en situation de handicap",
    )

    has_senior = any(pattern in text for pattern in senior_patterns)
    has_professional = any(pattern in text for pattern in professional_patterns)
    has_other = any(pattern in text for pattern in other_patterns)

    if has_senior and (has_professional or has_other):
        return "ainés et autres"
    if has_senior:
        return "ainés"
    if has_professional:
        return "professionnels"
    if has_other:
        return "autre"
    return "public non connu"


def nature_from_keywords(text: str) -> str:
    """Return the first fixed-priority keyword category, or an empty label."""
    text = normalise(text)
    for nature, keywords in NATURE_RULES:
        if any(keyword in text for keyword in keywords):
            return nature
    return ""


def final_path_from_llm_path(llm_path: str) -> Path:
    path = Path(llm_path)
    try:
        relative = path.relative_to(LLM_DATABASE)
    except ValueError as exc:
        raise ValueError(f"LLM JSON path is outside LLM_Tagged_Database: {path}") from exc
    return FINAL_DATABASE / relative


def load_completed_workbook(path: Path) -> tuple[dict[str, dict[str, object]], list[dict[str, object]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    required_sheets = {"Documents", "Annotations"}
    if not required_sheets.issubset(workbook.sheetnames):
        raise ValueError(f"Workbook must contain sheets: {sorted(required_sheets)}")

    document_rows = workbook["Documents"].iter_rows(values_only=True)
    document_headers = next(document_rows)
    document_indexes = {str(name): index for index, name in enumerate(document_headers)}
    required_document_columns = {"document_id", "json_file", "text_file"}
    if not required_document_columns.issubset(document_indexes):
        raise ValueError("Documents sheet is missing required columns.")

    documents: dict[str, dict[str, object]] = {}
    for row in document_rows:
        document_id = str(row[document_indexes["document_id"]] or "").strip()
        if document_id:
            documents[document_id] = {
                "json_file": row[document_indexes["json_file"]],
                "text_file": row[document_indexes["text_file"]],
            }

    annotation_rows = workbook["Annotations"].iter_rows(values_only=True)
    annotation_headers = next(annotation_rows)
    annotation_indexes = {str(name): index for index, name in enumerate(annotation_headers)}
    required_annotation_columns = {"document_id", "field", "gold_value"}
    if not required_annotation_columns.issubset(annotation_indexes):
        raise ValueError("Annotations sheet is missing required columns.")

    annotations: list[dict[str, object]] = []
    for row in annotation_rows:
        document_id = str(row[annotation_indexes["document_id"]] or "").strip()
        field = str(row[annotation_indexes["field"]] or "").strip()
        if document_id and field in FIELDS:
            annotations.append({
                "document_id": document_id,
                "field": field,
                "gold_value": row[annotation_indexes["gold_value"]],
            })
    return documents, annotations


def predict_documents(documents: dict[str, dict[str, object]]) -> dict[str, dict[str, object]]:
    predictions: dict[str, dict[str, object]] = {}
    for document_id, document in documents.items():
        final_path = final_path_from_llm_path(str(document["json_file"]))
        text_path = Path(str(document["text_file"]))
        if not final_path.exists():
            raise FileNotFoundError(f"Rule-based JSON is missing for {document_id}: {final_path}")
        if not text_path.exists():
            raise FileNotFoundError(f"Cleaned text is missing for {document_id}: {text_path}")

        record = json.loads(final_path.read_text(encoding="utf-8"))
        cleaned_text = text_path.read_text(encoding="utf-8", errors="ignore")[:MAX_CHARS]
        predictions[document_id] = {
            "public_vise": public_from_rule_output(record.get("public_vise")),
            "echelle": record.get("echelle", []),
            "nature_initiative": nature_from_keywords(cleaned_text),
            "raw_public_vise": record.get("public_vise", ""),
            "raw_echelle": record.get("echelle", []),
            "nature_input_characters": len(cleaned_text),
        }
    return predictions


def score(annotation_rows: list[dict[str, object]], predictions: dict[str, dict[str, object]]) -> list[dict[str, object]]:
    stats = defaultdict(lambda: Counter(tp=0, fp=0, fn=0, exact=0, n=0))
    label_stats = defaultdict(lambda: defaultdict(lambda: Counter(tp=0, fp=0, fn=0)))
    for row in annotation_rows:
        field = str(row["field"])
        gold = labels(row["gold_value"])
        predicted = labels(predictions[str(row["document_id"])][field])
        stat = stats[field]
        stat["tp"] += len(predicted & gold)
        stat["fp"] += len(predicted - gold)
        stat["fn"] += len(gold - predicted)
        stat["exact"] += int(predicted == gold)
        stat["n"] += 1
        for label in predicted | gold:
            label_stat = label_stats[field][label]
            if label in predicted and label in gold:
                label_stat["tp"] += 1
            elif label in predicted:
                label_stat["fp"] += 1
            else:
                label_stat["fn"] += 1

    results: list[dict[str, object]] = []
    for field in FIELDS:
        stat = stats[field]
        precision = safe_ratio(stat["tp"], stat["tp"] + stat["fp"])
        recall = safe_ratio(stat["tp"], stat["tp"] + stat["fn"])
        per_label_f1 = []
        for label_stat in label_stats[field].values():
            label_precision = safe_ratio(label_stat["tp"], label_stat["tp"] + label_stat["fp"])
            label_recall = safe_ratio(label_stat["tp"], label_stat["tp"] + label_stat["fn"])
            per_label_f1.append(f1(label_precision, label_recall))
        results.append({
            "method": "rule_based",
            "field": field,
            "n_annotated": stat["n"],
            "precision_micro": precision,
            "recall_micro": recall,
            "f1_micro": f1(precision, recall),
            "f1_macro": round(sum(per_label_f1) / len(per_label_f1), 3) if per_label_f1 else 0.0,
            "exact_match_rate": safe_ratio(stat["exact"], stat["n"]),
        })
    return results


def load_llm_results() -> dict[str, dict[str, str]]:
    path = OUTPUT_DIR / "llm_evaluation_results.csv"
    if not path.exists():
        raise FileNotFoundError(f"LLM result file not found: {path}")
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return {row["field"]: row for row in csv.DictReader(handle, delimiter=";") if row["field"] in FIELDS}


def write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    columns = sorted({column for row in rows for column in row})
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, delimiter=";")
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    input_path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_INPUT
    if not input_path.exists():
        raise FileNotFoundError(f"Annotation workbook not found: {input_path}")

    documents, annotations = load_completed_workbook(input_path)
    if len(documents) != 60:
        print(f"Warning: expected 60 documents, found {len(documents)}.")
    predictions = predict_documents(documents)
    baseline_results = score(annotations, predictions)
    write_csv(BASELINE_RESULTS, baseline_results)

    llm_results = load_llm_results()
    comparison_rows: list[dict[str, object]] = []
    for baseline in baseline_results:
        field = str(baseline["field"])
        llm = llm_results.get(field, {})
        comparison_rows.append({
            "field": field,
            "n_annotated": baseline["n_annotated"],
            "rule_based_precision": baseline["precision_micro"],
            "rule_based_recall": baseline["recall_micro"],
            "rule_based_f1": baseline["f1_micro"],
            "rule_based_exact_match": baseline["exact_match_rate"],
            "llm_precision": llm.get("precision_micro", ""),
            "llm_recall": llm.get("recall_micro", ""),
            "llm_f1": llm.get("f1_micro", ""),
            "llm_exact_match": llm.get("exact_match_rate", ""),
            "f1_difference_llm_minus_rule": round(float(llm.get("f1_micro", 0.0)) - float(baseline["f1_micro"]), 3),
        })
    write_csv(COMPARISON_RESULTS, comparison_rows)

    diagnostic_rows: list[dict[str, object]] = []
    gold_by_document = {(str(row["document_id"]), str(row["field"])): row["gold_value"] for row in annotations}
    for document_id in sorted(documents):
        prediction = predictions[document_id]
        for field in FIELDS:
            predicted_value = prediction[field]
            gold_value = gold_by_document[(document_id, field)]
            diagnostic_rows.append({
                "document_id": document_id,
                "field": field,
                "gold_value": gold_value,
                "rule_based_prediction": " | ".join(predicted_value) if isinstance(predicted_value, list) else predicted_value,
                "exact_match": labels(predicted_value) == labels(gold_value),
                "raw_rule_based_public_vise": prediction["raw_public_vise"] if field == "public_vise" else "",
                "raw_rule_based_echelle": " | ".join(prediction["raw_echelle"]) if field == "echelle" else "",
                "nature_input_characters": prediction["nature_input_characters"] if field == "nature_initiative" else "",
            })
    write_csv(DIAGNOSTICS, diagnostic_rows)

    print(f"Created {BASELINE_RESULTS}")
    print(f"Created {COMPARISON_RESULTS}")
    print(f"Created {DIAGNOSTICS}")
    for row in comparison_rows:
        print(
            f"{row['field']}: rule F1={row['rule_based_f1']}, "
            f"LLM F1={row['llm_f1']}, difference={row['f1_difference_llm_minus_rule']}"
        )


if __name__ == "__main__":
    main()
