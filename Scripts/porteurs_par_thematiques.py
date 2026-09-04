from pathlib import Path
import json
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


BASE_DIR = Path(__file__).resolve().parents[1]

LLM_DIR = BASE_DIR / "LLM_Tagged_Database"
OUTPUT_FILE = BASE_DIR / "comparatif_porteurs_par_thematique.xlsx"

FICHES_ACTIONS_KEYWORD = "Fiches_actions_database"


def load_json(path):
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def normalize_text(value):
    if value is None:
        return ""

    value = str(value).strip().lower()

    while "  " in value:
        value = value.replace("  ", " ")

    return value


def extract_values(value):
    values = []

    if value is None:
        return values

    if isinstance(value, list):
        for item in value:
            values.extend(extract_values(item))

    elif isinstance(value, dict):
        for item in value.values():
            values.extend(extract_values(item))

    else:
        cleaned = normalize_text(value)
        if cleaned:
            values.append(cleaned)

    return values


def is_fiche_action(json_path):
    return FICHES_ACTIONS_KEYWORD.lower() in str(json_path).lower()


def collect_porteurs_by_thematique(json_files):
    porteurs_par_thematique = defaultdict(set)

    stats = {
        "total_files": 0,
        "valid_files": 0,
        "without_thematique": 0,
        "without_porteur": 0,
        "errors": 0
    }

    for json_path in json_files:
        stats["total_files"] += 1

        try:
            data = load_json(json_path)

            if not isinstance(data, dict):
                continue

            stats["valid_files"] += 1

            thematiques = extract_values(data.get("thematique"))
            porteurs = extract_values(data.get("porteur_initiative"))

            if not thematiques:
                stats["without_thematique"] += 1
                continue

            if not porteurs:
                stats["without_porteur"] += 1
                continue

            for thematique in thematiques:
                for porteur in porteurs:
                    porteurs_par_thematique[thematique].add(porteur)

        except Exception as e:
            stats["errors"] += 1
            print(f"Erreur avec {json_path.name} : {e}")

    return porteurs_par_thematique, stats


def write_porteurs_sheet(ws, data):
    ws.append([
        "Thématique",
        "Porteur d'initiative"
    ])

    for thematique in sorted(data.keys()):
        for porteur in sorted(data[thematique]):
            ws.append([
                thematique,
                porteur
            ])


def write_comparatif_sheet(ws, fiches_actions_data, autres_fiches_data):
    ws.append([
        "Thématique",
        "Nb porteurs fiches actions",
        "Nb porteurs autres fiches",
        "Porteurs communs",
        "Porteurs seulement fiches actions",
        "Porteurs seulement autres fiches"
    ])

    all_thematiques = sorted(
        set(fiches_actions_data.keys()) | set(autres_fiches_data.keys())
    )

    for thematique in all_thematiques:
        porteurs_fa = fiches_actions_data.get(thematique, set())
        porteurs_autres = autres_fiches_data.get(thematique, set())

        communs = sorted(porteurs_fa & porteurs_autres)
        seulement_fa = sorted(porteurs_fa - porteurs_autres)
        seulement_autres = sorted(porteurs_autres - porteurs_fa)

        ws.append([
            thematique,
            len(porteurs_fa),
            len(porteurs_autres),
            "\n".join(communs),
            "\n".join(seulement_fa),
            "\n".join(seulement_autres)
        ])


def write_resume_sheet(wb, stats_fa, stats_autres, data_fa, data_autres):
    ws = wb.create_sheet(title="Résumé")

    ws.append([
        "Indicateur",
        "Fiches actions",
        "Autres fiches"
    ])

    ws.append([
        "Fichiers JSON trouvés",
        stats_fa["total_files"],
        stats_autres["total_files"]
    ])

    ws.append([
        "Fichiers JSON valides analysés",
        stats_fa["valid_files"],
        stats_autres["valid_files"]
    ])

    ws.append([
        "Thématiques différentes",
        len(data_fa),
        len(data_autres)
    ])

    ws.append([
        "Fichiers sans thématique",
        stats_fa["without_thematique"],
        stats_autres["without_thematique"]
    ])

    ws.append([
        "Fichiers sans porteur d'initiative",
        stats_fa["without_porteur"],
        stats_autres["without_porteur"]
    ])

    ws.append([
        "Erreurs",
        stats_fa["errors"],
        stats_autres["errors"]
    ])


def format_workbook(wb):
    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True)

        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

        for col_idx in range(1, ws.max_column + 1):
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = 35

        ws.freeze_panes = "A2"


def main():
    print("Recherche des fichiers JSON...")

    json_files = list(LLM_DIR.rglob("*.json"))
    print(f"{len(json_files)} fichiers JSON trouvés.")

    if not json_files:
        print(f"Dossier introuvable ou vide : {LLM_DIR}")
        return

    fiches_actions_files = [
        path for path in json_files
        if is_fiche_action(path)
    ]

    autres_fiches_files = [
        path for path in json_files
        if not is_fiche_action(path)
    ]

    print(f"Fiches actions : {len(fiches_actions_files)}")
    print(f"Autres fiches : {len(autres_fiches_files)}")

    fiches_actions_data, stats_fa = collect_porteurs_by_thematique(fiches_actions_files)
    autres_fiches_data, stats_autres = collect_porteurs_by_thematique(autres_fiches_files)

    wb = Workbook()

    ws_fa = wb.active
    ws_fa.title = "Fiches actions"
    write_porteurs_sheet(ws_fa, fiches_actions_data)

    ws_autres = wb.create_sheet(title="Autres fiches")
    write_porteurs_sheet(ws_autres, autres_fiches_data)

    ws_comparatif = wb.create_sheet(title="Comparatif")
    write_comparatif_sheet(ws_comparatif, fiches_actions_data, autres_fiches_data)

    write_resume_sheet(wb, stats_fa, stats_autres, fiches_actions_data, autres_fiches_data)

    format_workbook(wb)

    wb.save(OUTPUT_FILE)

    print(f"Excel créé : {OUTPUT_FILE}")


if __name__ == "__main__":
    main()