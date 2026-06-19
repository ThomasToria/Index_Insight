from pathlib import Path
import json
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


BASE_DIR = Path(__file__).resolve().parents[1]

LLM_DIR = BASE_DIR / "LLM_Tagged_Database"
OUTPUT_FILE = BASE_DIR / "possibilites_categories_sans_doublons.xlsx"

CATEGORIES = [
    "territoire",
    "echelle",
    "public_vise",
    "nature_initiative",
    "porteur_initiative",
    "date",
    "thematique",
    "source_site",
    "statut_action"
]


def load_json(path):
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def normalize_value(value):
    if value is None:
        return ""

    value = str(value).strip().lower()

    while "  " in value:
        value = value.replace("  ", " ")

    return value


def extract_values(value):
    values = []

    if value is None:
        return values

    if isinstance(value, list):
        for item in value:
            values.extend(extract_values(item))

    elif isinstance(value, dict):
        for item in value.values():
            values.extend(extract_values(item))

    else:
        clean_value = normalize_value(value)
        if clean_value:
            values.append(clean_value)

    return values


def main():
    print("Recherche des fichiers JSON...")

    json_files = list(LLM_DIR.rglob("*.json"))

    print(f"{len(json_files)} fichiers JSON trouvés.")

    if not json_files:
        print(f"Dossier introuvable ou vide : {LLM_DIR}")
        return

    values_by_category = defaultdict(set)
    total_valid_files = 0

    for json_path in json_files:
        try:
            data = load_json(json_path)

            if not isinstance(data, dict):
                print(f"Ignoré : {json_path.name}")
                continue

            total_valid_files += 1

            for category in CATEGORIES:
                if category in data:
                    extracted_values = extract_values(data.get(category))

                    for value in extracted_values:
                        values_by_category[category].add(value)

        except Exception as e:
            print(f"Erreur avec {json_path.name} : {e}")

    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "Résumé"

    ws_summary.append([
        "Catégorie",
        "Nombre de valeurs uniques"
    ])

    for category in CATEGORIES:
        ws_summary.append([
            category,
            len(values_by_category[category])
        ])

    for cell in ws_summary[1]:
        cell.font = Font(bold=True)

    ws_all = wb.create_sheet(title="Toutes les valeurs")

    col_idx = 1

    for category in CATEGORIES:
        ws_all.cell(row=1, column=col_idx, value=category)

        values = sorted(values_by_category[category])

        for row_idx, value in enumerate(values, start=2):
            ws_all.cell(row=row_idx, column=col_idx, value=value)

        col_idx += 1

    for cell in ws_all[1]:
        cell.font = Font(bold=True)

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

        for col_idx in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 35

    wb.save(OUTPUT_FILE)

    print(f"Excel créé : {OUTPUT_FILE}")
    print(f"Nombre de fichiers JSON valides analysés : {total_valid_files}")


if __name__ == "__main__":
    main()