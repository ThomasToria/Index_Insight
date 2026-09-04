"""Score a completed human-annotation workbook created by this project."""

from __future__ import annotations

import csv
import sys
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = PROJECT_ROOT / "Evaluation" / "llm_annotations_reviewer1.xlsx"
OUTPUT_FILE = PROJECT_ROOT / "Evaluation" / "llm_evaluation_results.csv"
SOURCE_OUTPUT_FILE = PROJECT_ROOT / "Evaluation" / "llm_evaluation_by_source.csv"
CATEGORICAL_FIELDS = {"public_vise", "echelle", "nature_initiative"}
ACCEPTED_JUDGMENTS = {"correct", "absence_correcte"}
SOURCE_DISPLAY_NAMES = {
    "Fiches actions": "Fiches actions SIAGE",
    "IRESP": "IReSP Autonomie",
    "VADA": "RFVAA",
}


def normalise(value: object) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKD", text.casefold())
    text = "".join(char for char in text if not unicodedata.combining(char))
    return " ".join(text.split())


def labels(value: object) -> set[str]:
    text = normalise(value)
    if not text or text == "absent":
        return set()
    return {item.strip() for item in text.split("|") if item.strip()}


def safe_ratio(numerator: int, denominator: int) -> float:
    return round(numerator / denominator, 3) if denominator else 0.0


def f1(precision: float, recall: float) -> float:
    return round(2 * precision * recall / (precision + recall), 3) if precision + recall else 0.0


def score_rows(annotation_rows: list[dict[str, object]]) -> list[dict[str, object]]:
    """Calcule les métriques globales ou pour un sous-ensemble de lignes."""
    categorical = defaultdict(lambda: Counter(tp=0, fp=0, fn=0, exact=0, n=0, missing=0))
    categorical_labels = defaultdict(lambda: defaultdict(lambda: Counter(tp=0, fp=0, fn=0)))
    judgments = defaultdict(Counter)

    for row in annotation_rows:
        field = str(row["field"] or "").strip()
        judgment = normalise(row["judgment"])
        if not field or not judgment:
            continue
        if field in CATEGORICAL_FIELDS:
            gold_cell = row["gold_value"]
            if not str(gold_cell or "").strip():
                categorical[field]["missing"] += 1
                continue
            predicted = labels(row["predicted_value"])
            gold = labels(gold_cell)
            stat = categorical[field]
            stat["tp"] += len(predicted & gold)
            stat["fp"] += len(predicted - gold)
            stat["fn"] += len(gold - predicted)
            stat["exact"] += int(predicted == gold)
            stat["n"] += 1
            for label in predicted | gold:
                label_stat = categorical_labels[field][label]
                if label in predicted and label in gold:
                    label_stat["tp"] += 1
                elif label in predicted:
                    label_stat["fp"] += 1
                else:
                    label_stat["fn"] += 1
        else:
            judgments[field][judgment] += 1

    output_rows: list[dict[str, object]] = []
    for field, stat in sorted(categorical.items()):
        precision = safe_ratio(stat["tp"], stat["tp"] + stat["fp"])
        recall = safe_ratio(stat["tp"], stat["tp"] + stat["fn"])
        per_label_f1 = []
        for label_stat in categorical_labels[field].values():
            label_precision = safe_ratio(label_stat["tp"], label_stat["tp"] + label_stat["fp"])
            label_recall = safe_ratio(label_stat["tp"], label_stat["tp"] + label_stat["fn"])
            per_label_f1.append(f1(label_precision, label_recall))
        output_rows.append({
            "field": field,
            "metric_type": "categorical",
            "n_annotated": stat["n"],
            "precision_micro": precision,
            "recall_micro": recall,
            "f1_micro": f1(precision, recall),
            "f1_macro": round(sum(per_label_f1) / len(per_label_f1), 3) if per_label_f1 else 0.0,
            "exact_match_rate": safe_ratio(stat["exact"], stat["n"]),
            "missing_gold_values": stat["missing"],
        })

    for field, stat in sorted(judgments.items()):
        total = sum(stat.values())
        evaluable = total - stat["non_evaluable"]
        accepted = sum(stat[item] for item in ACCEPTED_JUDGMENTS)
        output_rows.append({
            "field": field,
            "metric_type": "human_judgment",
            "n_annotated": total,
            "n_evaluable": evaluable,
            "accepted_rate_evaluable": safe_ratio(accepted, evaluable),
            "accepted_rate_full_sample": safe_ratio(accepted, total),
            "correct": stat["correct"],
            "partiel": stat["partiel"],
            "incorrect": stat["incorrect"],
            "absence_correcte": stat["absence_correcte"],
            "non_evaluable": stat["non_evaluable"],
        })
    return output_rows


def write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    """Écrit un tableau de résultats, même si certains champs sont spécifiques."""
    if not rows:
        raise ValueError("No completed annotations found.")
    columns = sorted({column for row in rows for column in row})
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, delimiter=";")
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    input_file = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_INPUT
    if not input_file.exists():
        raise FileNotFoundError(
            f"Annotation workbook not found: {input_file}. Copy the template and complete it first."
        )

    workbook = load_workbook(input_file, read_only=True, data_only=True)
    if "Annotations" not in workbook.sheetnames:
        raise ValueError("Workbook must contain an 'Annotations' sheet.")
    rows = workbook["Annotations"].iter_rows(values_only=True)
    headers = next(rows)
    indexes = {name: position for position, name in enumerate(headers)}
    required = {"source", "field", "predicted_value", "gold_value", "judgment"}
    if not required.issubset(indexes):
        raise ValueError(f"Missing required columns: {sorted(required - set(indexes))}")

    annotations: list[dict[str, object]] = []
    for row in rows:
        annotations.append({
            "source": str(row[indexes["source"]] or "").strip(),
            "field": row[indexes["field"]],
            "predicted_value": row[indexes["predicted_value"]],
            "gold_value": row[indexes["gold_value"]],
            "judgment": row[indexes["judgment"]],
        })

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    output_rows = score_rows(annotations)
    write_csv(OUTPUT_FILE, output_rows)
    print(f"Created {OUTPUT_FILE}")

    source_rows: list[dict[str, object]] = []
    for source in sorted({str(item["source"]) for item in annotations if item["source"]}):
        source_annotations = [item for item in annotations if item["source"] == source]
        for result in score_rows(source_annotations):
            source_rows.append({
                "source": SOURCE_DISPLAY_NAMES.get(source, source),
                **result,
            })
    write_csv(SOURCE_OUTPUT_FILE, source_rows)
    print(f"Created {SOURCE_OUTPUT_FILE}")


if __name__ == "__main__":
    main()
