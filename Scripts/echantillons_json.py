from pathlib import Path
import json
import random
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


BASE_DIR = Path(__file__).resolve().parents[1]

LLM_DIR = BASE_DIR / "LLM_Tagged_Database"
CLEANED_DIR = BASE_DIR / "Cleaned_Database"
OUTPUT_FILE = BASE_DIR / "echantillon_llm_tagged.xlsx"

NB_FICHIERS = 10
RANDOM_SEED = None

CATEGORIES = [
    "titre",
    "territoire",
    "echelle",
    "public_vise",
    "description_generale",
    "contexte",
    "problematique",
    "solution_envisagee",
    "objectifs",
    "porteur_initiative",
    "date",
    "thematique",
    "source_site"
]


def load_json(path):
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def value_to_text(value):
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, indent=2)

    if value is None:
        return ""

    return str(value)


def find_matching_txt(json_path):
    stem = json_path.stem

    for txt_path in CLEANED_DIR.rglob("*.txt"):
        if txt_path.stem == stem:
            return txt_path

    return None


def read_text(path):
    if path is None:
        return "Fichier texte correspondant introuvable dans Cleaned_Database."

    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def main():
    print("Recherche des fichiers JSON...")

    if RANDOM_SEED is not None:
        random.seed(RANDOM_SEED)

    json_files = list(LLM_DIR.rglob("*.json"))

    print(f"{len(json_files)} fichiers JSON trouvés.")

    if len(json_files) == 0:
        print(f"Dossier introuvable ou vide : {LLM_DIR}")
        return

    selected_files = random.sample(
        json_files,
        min(NB_FICHIERS, len(json_files))
    )

    print("Fichiers sélectionnés :")

    for f in selected_files:
        print(f" - {f.name}")

    data_by_file = {}

    for json_path in selected_files:
        try:
            data = load_json(json_path)

            if not isinstance(data, dict):
                print(f"Ignoré : {json_path.name}")
                continue

            data_by_file[json_path] = data

        except Exception as e:
            print(f"Erreur avec {json_path.name} : {e}")

    if len(data_by_file) == 0:
        print("Aucun JSON valide exploitable.")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "Echantillon"

    ws.cell(row=1, column=1, value="Catégorie")

    for col_idx, json_path in enumerate(data_by_file.keys(), start=2):
        ws.cell(row=1, column=col_idx, value=json_path.name)

    for row_idx, category in enumerate(CATEGORIES, start=2):
        ws.cell(row=row_idx, column=1, value=category)

        for col_idx, (json_path, data) in enumerate(data_by_file.items(), start=2):
            ws.cell(
                row=row_idx,
                column=col_idx,
                value=value_to_text(data.get(category, ""))
            )

    text_row = len(CATEGORIES) + 2

    ws.cell(row=text_row, column=1, value="texte_complet")

    print("Ajout des textes complets...")

    for col_idx, json_path in enumerate(data_by_file.keys(), start=2):
        txt_path = find_matching_txt(json_path)
        full_text = read_text(txt_path)

        if txt_path is None:
            print(f"Texte introuvable pour : {json_path.name}")
        else:
            print(f"Texte trouvé pour : {json_path.name}")

        ws.cell(
            row=text_row,
            column=col_idx,
            value=full_text
        )

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                wrap_text=True,
                vertical="top"
            )

    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row_idx in range(2, text_row + 1):
        ws.cell(row=row_idx, column=1).font = Font(bold=True)

    ws.column_dimensions["A"].width = 35

    for col_idx in range(2, len(data_by_file) + 2):
        ws.column_dimensions[get_column_letter(col_idx)].width = 45

    ws.freeze_panes = "B2"

    print("Sauvegarde du fichier Excel...")

    wb.save(OUTPUT_FILE)

    print(f"Excel créé : {OUTPUT_FILE}")
    print(f"Nombre de fichiers sélectionnés : {len(data_by_file)}")


if __name__ == "__main__":
    main()