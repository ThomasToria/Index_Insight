"""Create a reproducible, stratified human-annotation workbook for LLM outputs."""

from __future__ import annotations

import json
import random
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[1]
LLM_DIR = PROJECT_ROOT / "LLM_Tagged_Database"
CLEANED_DIR = PROJECT_ROOT / "Cleaned_Database"
OUTPUT_FILE = PROJECT_ROOT / "Evaluation" / "llm_annotation_template.xlsx"

SEED = 20260717
DOCUMENTS_PER_SOURCE = 10
FIELDS = [
    "territoire",
    "echelle",
    "public_vise",
    "nature_initiative",
    "porteur_initiative",
    "date",
    "thematique",
]


def read_json(path: Path) -> dict[str, Any] | None:
    try:
        with path.open("r", encoding="utf-8") as handle:
            value = json.load(handle)
        return value if isinstance(value, dict) else None
    except (OSError, json.JSONDecodeError):
        return None


def as_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return " | ".join(as_text(item) for item in value)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False)
    return str(value).strip()


def prediction(record: dict[str, Any], field: str) -> str:
    """Read both the current flat and older nested output schemas."""
    if field in record:
        return as_text(record[field])
    description = record.get("description")
    if isinstance(description, dict) and field in description:
        return as_text(description[field])
    return ""


def text_path_for(json_path: Path, record: dict[str, Any]) -> Path | None:
    declared = record.get("fichier_source")
    if isinstance(declared, str) and declared:
        candidate = Path(declared)
        if candidate.exists():
            return candidate

    filename = record.get("nom_fichier")
    if not isinstance(filename, str) or not filename:
        filename = f"{json_path.stem}.txt"

    matches = list(CLEANED_DIR.rglob(filename))
    return matches[0] if len(matches) == 1 else None


def load_candidates() -> dict[str, list[tuple[Path, Path, dict[str, Any]]]]:
    candidates: dict[str, list[tuple[Path, Path, dict[str, Any]]]] = {}
    for json_path in sorted(LLM_DIR.rglob("*.json")):
        record = read_json(json_path)
        if record is None:
            continue
        source_text_path = text_path_for(json_path, record)
        if source_text_path is None:
            continue
        source = as_text(record.get("source_site")) or json_path.parent.name
        candidates.setdefault(source, []).append((json_path, source_text_path, record))
    return candidates


def source_text(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8", errors="replace")
    except OSError:
        return ""


def format_sheet(sheet) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def main() -> None:
    candidates = load_candidates()
    if not candidates:
        raise RuntimeError("No valid LLM output matched to a cleaned source text.")

    rng = random.Random(SEED)
    selected: list[tuple[str, Path, Path, dict[str, Any]]] = []
    for source in sorted(candidates, key=str.casefold):
        records = candidates[source]
        if len(records) < DOCUMENTS_PER_SOURCE:
            raise RuntimeError(f"Source {source}: only {len(records)} usable documents.")
        for json_path, text_path, record in rng.sample(records, DOCUMENTS_PER_SOURCE):
            selected.append((source, json_path, text_path, record))

    workbook = Workbook()
    instructions = workbook.active
    instructions.title = "Instructions"
    instructions.append(["Evaluation LLM — annotation humaine"])
    instructions.append(["Échantillon aléatoire stratifié, graine", SEED])
    instructions.append(["Documents par source", DOCUMENTS_PER_SOURCE])
    instructions.append(["Consultez Evaluation/annotation_guide.md avant de commencer."])
    instructions.column_dimensions["A"].width = 80

    documents_sheet = workbook.create_sheet("Documents")
    documents_sheet.append([
        "document_id", "source", "json_file", "text_file", "source_text"
    ])

    annotations_sheet = workbook.create_sheet("Annotations")
    annotations_sheet.append([
        "document_id", "source", "field", "predicted_value", "gold_value",
        "judgment", "notes"
    ])

    for number, (source, json_path, text_path, record) in enumerate(selected, start=1):
        document_id = f"D{number:03d}"
        documents_sheet.append([
            document_id, source, str(json_path), str(text_path), source_text(text_path)
        ])
        for field in FIELDS:
            annotations_sheet.append([
                document_id, source, field, prediction(record, field), "", "", ""
            ])

    format_sheet(documents_sheet)
    format_sheet(annotations_sheet)
    documents_sheet.column_dimensions["A"].width = 13
    documents_sheet.column_dimensions["B"].width = 22
    documents_sheet.column_dimensions["C"].width = 55
    documents_sheet.column_dimensions["D"].width = 55
    documents_sheet.column_dimensions["E"].width = 100
    annotations_sheet.column_dimensions["A"].width = 13
    annotations_sheet.column_dimensions["B"].width = 22
    annotations_sheet.column_dimensions["C"].width = 24
    annotations_sheet.column_dimensions["D"].width = 45
    annotations_sheet.column_dimensions["E"].width = 45
    annotations_sheet.column_dimensions["F"].width = 20
    annotations_sheet.column_dimensions["G"].width = 60

    input_fill = PatternFill("solid", fgColor="FFF2CC")
    for row in annotations_sheet.iter_rows(min_row=2, min_col=5, max_col=7):
        for cell in row:
            cell.fill = input_fill

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT_FILE)
    print(f"Created {OUTPUT_FILE}")
    print(f"Selected {len(selected)} documents across {len(candidates)} sources.")


if __name__ == "__main__":
    main()
